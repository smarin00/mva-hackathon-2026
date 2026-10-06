import sys
import docx

d = docx.Document(sys.argv[1])
for p in d.paragraphs:
    if p.text.strip():
        print(p.text)
for t in d.tables:
    for r in t.rows:
        print(" | ".join(c.text for c in r.cells))
